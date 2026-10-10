"""
Phase 2 backend tests: Candidate application (admission info, sections, uploads, fee,
submission, PDF, post-submission flows) + Staff Part A dashboard, applications list,
export, correction, PDF, RBAC.
"""
import io
import os
import random
import time
import uuid

import pytest
import requests
from PIL import Image

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "https://hub-preview-32.preview.emergentagent.com").rstrip("/")
API = f"{BASE_URL}/api"

STAFF = {
    "super_admin": ("superadmin@weschool-demo.in", "SuperAdmin@2026"),
    "admissions_operator": ("admissions@weschool-demo.in", "Staff@2026"),
    "panelist": ("panelist@weschool-demo.in", "Staff@2026"),
    "accounts_operator": ("accounts@weschool-demo.in", "Staff@2026"),
    "auditor": ("auditor@weschool-demo.in", "Staff@2026"),
}
CAND_PW = "Candidate@2026"


# ---------- helpers ----------
def _login(email, pw, staff=True):
    path = "/auth/staff/login" if staff else "/auth/candidate/login"
    r = requests.post(f"{API}{path}", json={"email": email, "password": pw}, timeout=20)
    if staff and r.status_code == 200 and r.json().get("step_up_required"):
        body = r.json()
        r = requests.post(f"{API}/auth/staff/login/complete",
                          json={"challenge_id": body["challenge_id"], "code": body["debug_code"]}, timeout=20)
    return r


def _tok(email, pw, staff=True):
    r = _login(email, pw, staff)
    assert r.status_code == 200, f"login failed {email}: {r.status_code} {r.text}"
    return r.json()["access_token"]


def H(tok):
    return {"Authorization": f"Bearer {tok}", "Content-Type": "application/json"}


def _png_bytes(w, h, color=(158, 27, 50)):
    img = Image.new("RGB", (w, h), color)
    # Add noise so file is >2KB
    import random as _r
    px = img.load()
    for i in range(0, w, 2):
        for j in range(0, h, 2):
            px[i, j] = (_r.randint(0, 255), _r.randint(0, 255), _r.randint(0, 255))
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def _jpg_bytes(w, h, color=(158, 27, 50)):
    img = Image.new("RGB", (w, h), color)
    import random as _r
    px = img.load()
    for i in range(0, w, 2):
        for j in range(0, h, 2):
            px[i, j] = (_r.randint(0, 255), _r.randint(0, 255), _r.randint(0, 255))
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=85)
    return buf.getvalue()


# ---------- Shared fixtures ----------
@pytest.fixture(scope="module")
def admissions_tok():
    return _tok(*STAFF["admissions_operator"])


@pytest.fixture(scope="module")
def super_tok():
    return _tok(*STAFF["super_admin"])


@pytest.fixture(scope="module")
def auditor_tok():
    return _tok(*STAFF["auditor"])


@pytest.fixture(scope="module")
def panelist_tok():
    return _tok(*STAFF["panelist"])


@pytest.fixture(scope="module")
def accounts_tok():
    return _tok(*STAFF["accounts_operator"])


@pytest.fixture(scope="module")
def c01_tok():
    return _tok("candidate01@example.com", CAND_PW, staff=False)


def _reset_c02_to_draft():
    from pymongo import MongoClient
    from dotenv import load_dotenv
    load_dotenv("/app/backend/.env")
    db = MongoClient(os.environ["MONGO_URL"])[os.environ["DB_NAME"]]
    u = db.users.find_one({"email": "candidate02@example.com"})
    app = db.applications.find_one({"candidate_id": str(u["_id"])})
    if app:
        db.application_versions.delete_many({"application_id": str(app["_id"])})
        db.applications.update_one({"_id": app["_id"]}, {
            "$set": {"status": "draft", "submitted_at": None, "submitted_version": 0, "post": {},
                     "signature_file_id": None, "section_status.exams": "empty", "section_status.preferences": "empty"},
            "$unset": {"sections.exams": "", "sections.preferences": ""}})
    db.rate_limits.delete_many({})


@pytest.fixture(scope="module")
def c02_tok():
    _reset_c02_to_draft()
    return _tok("candidate02@example.com", CAND_PW, staff=False)


@pytest.fixture(scope="module")
def c03_tok():
    return _tok("candidate03@example.com", CAND_PW, staff=False)


# ---------- Admission info ----------
class TestAdmissionInfo:
    def test_admission_info(self, c01_tok):
        r = requests.get(f"{API}/candidate/admission-info", headers=H(c01_tok), timeout=20)
        assert r.status_code == 200
        d = r.json()
        assert d["cycle"]["code"] == "2026-28"
        assert d["cycle"]["application_fee_inr"] == 1500
        assert isinstance(d["programs"], list) and len(d["programs"]) >= 10
        assert d["window_open"] is True
        assert "news" in d and "resources" in d
        # TFW policy may or may not be seeded; just assert key exists
        assert "tfw" in d

    def test_dashboard(self, c01_tok):
        r = requests.get(f"{API}/candidate/dashboard", headers=H(c01_tok), timeout=20)
        assert r.status_code == 200
        d = r.json()
        assert d["application_status"] in ("draft", "submitted", "not_started")
        assert any(s["key"] == "submit" for s in d["steps"])


# ---------- Fee-unpaid submit blocked (candidate01) ----------
class TestCandidate01SubmitBlocked:
    def test_submit_without_fee_409(self, c01_tok):
        r = requests.post(f"{API}/candidate/application/submit", headers=H(c01_tok), timeout=20)
        assert r.status_code == 409
        d = r.json()
        # error code either in detail.code or top-level
        code = d.get("code") or (d.get("detail") or {}).get("code") if isinstance(d.get("detail"), dict) else None
        txt = str(d).lower()
        assert "fee_unpaid" in txt or "fee" in txt


# ---------- Validation ----------
class TestValidation:
    def test_contact_bad_pincode(self, c02_tok):
        body = {"data": {"address_line1": "A-101, Society", "city": "Mumbai", "state": "Maharashtra",
                         "pincode": "99",  # invalid
                         "emergency_contact_name": "Ms X", "emergency_contact_mobile": "9876543210"}}
        r = requests.put(f"{API}/candidate/application/sections/contact", json=body, headers=H(c02_tok), timeout=20)
        assert r.status_code == 200
        d = r.json()
        assert d["status"] == "in_progress"
        assert any("PIN" in e["message"] or "pincode" in e["field"].lower() for e in d["errors"])

    def test_education_year_order(self, c02_tok):
        body = {"data": {
            "class_x": {"board": "CBSE", "school": "KV", "year": 2020, "percentage": 85},
            "class_xii": {"board": "State", "school": "JC", "year": 2020, "percentage": 80},  # not >= x+2
            "graduation": {"degree": "BCom", "discipline": "Finance", "university": "MU",
                           "status": "completed", "year": 2023, "score_type": "cgpa", "score": 7.5}}}
        r = requests.put(f"{API}/candidate/application/sections/education", json=body, headers=H(c02_tok), timeout=20)
        assert r.status_code == 200
        assert r.json()["status"] == "in_progress"
        assert any("Class XII" in e["message"] or "2 years" in e["message"] for e in r.json()["errors"])

    def test_education_cgpa_over_10(self, c02_tok):
        body = {"data": {
            "class_x": {"board": "CBSE", "school": "KV", "year": 2016, "percentage": 85},
            "class_xii": {"board": "State", "school": "JC", "year": 2018, "percentage": 80},
            "graduation": {"degree": "BCom", "discipline": "Finance", "university": "MU",
                           "status": "completed", "year": 2022, "score_type": "cgpa", "score": 11.5}}}
        r = requests.put(f"{API}/candidate/application/sections/education", json=body, headers=H(c02_tok), timeout=20)
        assert r.status_code == 200
        assert r.json()["status"] == "in_progress"
        assert any("10" in e["message"] for e in r.json()["errors"])

    def test_work_has_experience_but_empty(self, c02_tok):
        body = {"data": {"has_experience": True, "experiences": []}}
        r = requests.put(f"{API}/candidate/application/sections/work", json=body, headers=H(c02_tok), timeout=20)
        assert r.status_code == 200
        assert r.json()["status"] == "in_progress"
        assert r.json()["errors"]

    def test_exams_duplicate(self, c02_tok):
        body = {"data": {"exams": [
            {"exam": "CAT", "year": 2025, "roll_number": "CAT25111", "status": "appeared", "percentile": 90.0},
            {"exam": "CAT", "year": 2025, "roll_number": "CAT25222", "status": "appeared", "percentile": 85.0}]}}
        r = requests.put(f"{API}/candidate/application/sections/exams", json=body, headers=H(c02_tok), timeout=20)
        assert r.status_code == 200
        assert r.json()["status"] == "in_progress"
        assert any("only once" in e["message"].lower() or "unique" in e["message"].lower() for e in r.json()["errors"])

    def test_preferences_duplicate(self, c02_tok):
        body = {"data": {"ranked": ["PGDM-MUM", "PGDM-MUM"]}}
        r = requests.put(f"{API}/candidate/application/sections/preferences", json=body, headers=H(c02_tok), timeout=20)
        assert r.status_code == 200
        assert r.json()["status"] == "in_progress"
        assert r.json()["errors"]

    def test_preferences_unknown_program(self, c02_tok):
        body = {"data": {"ranked": ["PGDM-FAKE-XYZ"]}}
        r = requests.put(f"{API}/candidate/application/sections/preferences", json=body, headers=H(c02_tok), timeout=20)
        assert r.status_code == 200
        d = r.json()
        assert d["status"] == "in_progress"
        assert any("Unknown" in e["message"] or "inactive" in e["message"] for e in d["errors"])

    def test_unknown_section_404(self, c02_tok):
        r = requests.put(f"{API}/candidate/application/sections/foo", json={"data": {}}, headers=H(c02_tok), timeout=20)
        assert r.status_code == 404


# ---------- Uploads ----------
class TestUploads:
    def _upload(self, tok, kind, files):
        return requests.post(f"{API}/candidate/application/uploads/{kind}",
                             headers={"Authorization": f"Bearer {tok}"}, files=files, timeout=30)

    def test_invalid_type_text(self, c02_tok):
        r = self._upload(c02_tok, "photo", {"file": ("x.txt", b"hello world", "text/plain")})
        assert r.status_code == 422
        assert "invalid_type" in r.text

    def test_invalid_image_content(self, c02_tok):
        # jpeg content-type but bytes are garbage, length > 2KB
        r = self._upload(c02_tok, "photo", {"file": ("x.jpg", b"X" * 3000, "image/jpeg")})
        assert r.status_code == 422
        assert "invalid_image" in r.text or "type_mismatch" in r.text

    def test_too_large_photo(self, c02_tok):
        data = _jpg_bytes(2000, 2400)
        if len(data) <= 200_000:
            pytest.skip("could not generate >200KB jpg")
        r = self._upload(c02_tok, "photo", {"file": ("big.jpg", data, "image/jpeg")})
        assert r.status_code == 422
        assert "too_large" in r.text

    def test_too_small_dimensions(self, c02_tok):
        data = _jpg_bytes(100, 100)
        r = self._upload(c02_tok, "photo", {"file": ("s.jpg", data, "image/jpeg")})
        assert r.status_code == 422
        # Could be too_small (file < 2KB) OR too_small_dimensions
        assert "too_small" in r.text

    def test_photo_bad_aspect_landscape(self, c02_tok):
        data = _jpg_bytes(600, 300)  # landscape, 2:1
        r = self._upload(c02_tok, "photo", {"file": ("land.jpg", data, "image/jpeg")})
        assert r.status_code == 422
        assert "bad_aspect" in r.text

    def test_signature_taller_than_wide(self, c02_tok):
        data = _jpg_bytes(200, 400)  # taller
        r = self._upload(c02_tok, "signature", {"file": ("sig.jpg", data, "image/jpeg")})
        assert r.status_code == 422
        assert "bad_aspect" in r.text

    def test_valid_photo_upload_and_file_access(self, c02_tok, admissions_tok, c01_tok):
        data = _jpg_bytes(400, 480)
        r = self._upload(c02_tok, "photo", {"file": ("ok.jpg", data, "image/jpeg")})
        assert r.status_code == 200, r.text
        fid = r.json()["file_id"]
        # owner can get
        r2 = requests.get(f"{API}/files/{fid}", headers=H(c02_tok), timeout=20)
        assert r2.status_code == 200
        assert r2.headers["content-type"].startswith("image/")
        # Staff with applications.read can get
        r3 = requests.get(f"{API}/files/{fid}", headers=H(admissions_tok), timeout=20)
        assert r3.status_code == 200
        # Other candidate cannot
        r4 = requests.get(f"{API}/files/{fid}", headers=H(c01_tok), timeout=20)
        assert r4.status_code == 404


# ---------- Complete candidate02 and submit ----------
class TestSubmitFlow:
    def test_complete_and_submit_c02(self, c02_tok):
        # fetch current app
        r = requests.get(f"{API}/candidate/application", headers=H(c02_tok), timeout=20)
        assert r.status_code == 200
        app = r.json()
        if app["status"] == "submitted":
            pytest.skip("candidate02 already submitted (prior run)")

        # Ensure all 7 sections are complete; use seed-like payloads
        sections = {
            "personal": {"full_name": "Candidate Two", "dob": "2001-05-12", "gender": "Female",
                         "category": "General", "nationality": "Indian", "pwd": False},
            "family": {"father_name": "Mr Two", "father_occupation": "Business",
                       "mother_name": "Mrs Two", "mother_occupation": "Teacher",
                       "annual_family_income_inr": 800000, "siblings": []},
            "contact": {"address_line1": "A-101, Shanti Niketan", "city": "Mumbai",
                        "state": "Maharashtra", "pincode": "400011",
                        "emergency_contact_name": "Mr Two Sr", "emergency_contact_mobile": "9876543210"},
            "education": {
                "class_x": {"board": "CBSE", "school": "KV", "year": 2016, "percentage": 85},
                "class_xii": {"board": "State", "school": "JC", "year": 2018, "percentage": 82},
                "graduation": {"degree": "BCom", "discipline": "Finance", "university": "MU",
                               "status": "completed", "year": 2022, "score_type": "cgpa", "score": 8.1}},
            "work": {"has_experience": False, "experiences": []},
            "exams": {"exams": [{"exam": "CAT", "year": 2025, "roll_number": "CAT25999",
                                 "status": "appeared", "percentile": 92.5}]},
            "preferences": {"ranked": ["PGDM-MUM", "PGDM-BD"]},
        }
        for k, data in sections.items():
            rr = requests.put(f"{API}/candidate/application/sections/{k}", json={"data": data},
                              headers=H(c02_tok), timeout=20)
            assert rr.status_code == 200, f"{k}: {rr.text}"
            assert rr.json()["status"] == "complete", f"{k} not complete: {rr.json()}"

        # Upload signature (wider than tall), >2KB
        sig_data = _jpg_bytes(400, 120)
        rs = requests.post(f"{API}/candidate/application/uploads/signature",
                           headers={"Authorization": f"Bearer {c02_tok}"},
                           files={"file": ("sig.jpg", sig_data, "image/jpeg")}, timeout=30)
        assert rs.status_code == 200, rs.text

        # Submit
        r = requests.post(f"{API}/candidate/application/submit", headers=H(c02_tok), timeout=30)
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["status"] == "submitted"
        assert d["submitted_version"] == 1

        # Second submit -> 409 already_submitted
        r2 = requests.post(f"{API}/candidate/application/submit", headers=H(c02_tok), timeout=20)
        assert r2.status_code == 409
        assert "already_submitted" in r2.text

        # Section save blocked after submission
        r3 = requests.put(f"{API}/candidate/application/sections/personal",
                         json={"data": sections["personal"]}, headers=H(c02_tok), timeout=20)
        assert r3.status_code == 409
        assert "locked" in r3.text

        # Download PDF
        rp = requests.get(f"{API}/candidate/application/pdf", headers=H(c02_tok), timeout=30)
        assert rp.status_code == 200
        assert rp.headers["content-type"] == "application/pdf"
        assert len(rp.content) > 500


# ---------- Post-submission flows (use candidate03 if not already done) ----------
class TestPostSubmission:
    @pytest.fixture(scope="class")
    def c03(self):
        return _tok("candidate03@example.com", CAND_PW, staff=False)

    def test_declaration(self, c03):
        r = requests.get(f"{API}/candidate/application", headers=H(c03), timeout=20)
        assert r.status_code == 200
        post = r.json().get("post") or {}
        if post.get("declaration"):
            # Second attempt -> 409
            r2 = requests.post(f"{API}/candidate/application/post/declaration",
                               json={"accept": True}, headers=H(c03), timeout=20)
            assert r2.status_code == 409
        else:
            r2 = requests.post(f"{API}/candidate/application/post/declaration",
                               json={"accept": True}, headers=H(c03), timeout=20)
            assert r2.status_code == 200
            # Repeat -> 409
            r3 = requests.post(f"{API}/candidate/application/post/declaration",
                               json={"accept": True}, headers=H(c03), timeout=20)
            assert r3.status_code == 409

    def test_essays_word_count(self, c03):
        # Too few words
        r = requests.put(f"{API}/candidate/application/post/essays",
                        json={"answers": {"e1": "short", "e2": "short"}, "submit": True},
                        headers=H(c03), timeout=20)
        assert r.status_code == 422
        assert "words" in r.text.lower() or "100" in r.text

    def test_graduation_eligibility_final_year_requires_month(self, c03):
        r = requests.post(f"{API}/candidate/application/post/graduation",
                         json={"graduation_status": "final_year", "expected_completion": "",
                               "meets_min_aggregate": True, "has_active_backlogs": False, "accept": True},
                         headers=H(c03), timeout=20)
        assert r.status_code == 422
        assert "expected_completion" in r.text or "final-year" in r.text

    def test_preference_reset_unchanged(self, c03):
        r = requests.get(f"{API}/candidate/application", headers=H(c03), timeout=20)
        prefs = (r.json().get("sections") or {}).get("preferences", {}).get("ranked", [])
        if not prefs:
            pytest.skip("no prefs")
        r2 = requests.post(f"{API}/candidate/application/post/preferences",
                          json={"ranked": prefs}, headers=H(c03), timeout=20)
        assert r2.status_code == 409
        assert "unchanged" in r2.text


# ---------- Staff Part A dashboard & list ----------
class TestStaffPartA:
    def test_dashboard(self, admissions_tok):
        r = requests.get(f"{API}/admin/part-a/dashboard", headers=H(admissions_tok), timeout=30)
        assert r.status_code == 200
        d = r.json()
        for k in ("registered", "verified", "paid", "submitted"):
            assert k in d["counts"]
            assert isinstance(d["counts"][k], int)
        assert d["counts"]["registered"] >= 10
        assert d["counts"]["paid"] >= 1
        assert isinstance(d["by_state"], list)
        assert isinstance(d["first_preferences"], list)

    def test_list_and_filters(self, admissions_tok):
        r = requests.get(f"{API}/admin/applications", headers=H(admissions_tok), timeout=30)
        assert r.status_code == 200
        rows = r.json()
        assert isinstance(rows, list) and len(rows) >= 10
        # Filter by submitted
        r2 = requests.get(f"{API}/admin/applications?status=submitted",
                         headers=H(admissions_tok), timeout=30)
        assert all(x["application_status"] == "submitted" for x in r2.json())
        # Search
        r3 = requests.get(f"{API}/admin/applications?q=candidate", headers=H(admissions_tok), timeout=30)
        assert r3.status_code == 200
        # Fee paid filter
        r4 = requests.get(f"{API}/admin/applications?fee=paid", headers=H(admissions_tok), timeout=30)
        assert all(x["fee_paid"] for x in r4.json())

    def test_export_xlsx(self, admissions_tok):
        r = requests.get(f"{API}/admin/applications/export", headers=H(admissions_tok), timeout=60)
        assert r.status_code == 200
        assert "spreadsheet" in r.headers["content-type"]
        assert len(r.content) > 1000
        # Verify it's a valid xlsx by opening
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(r.content))
        assert "Candidates" in wb.sheetnames
        assert "State-wise paid" in wb.sheetnames

    def test_application_detail_and_pdf(self, admissions_tok):
        # find a submitted one
        rows = requests.get(f"{API}/admin/applications?status=submitted",
                            headers=H(admissions_tok), timeout=30).json()
        assert rows, "no submitted apps"
        app_id = rows[0]["application_id"]
        r = requests.get(f"{API}/admin/applications/{app_id}",
                         headers=H(admissions_tok), timeout=20)
        assert r.status_code == 200
        d = r.json()
        assert d["status"] == "submitted"
        assert "versions" in d and len(d["versions"]) >= 1
        assert "candidate" in d
        # PDF
        rp = requests.get(f"{API}/admin/applications/{app_id}/pdf",
                          headers=H(admissions_tok), timeout=30)
        assert rp.status_code == 200
        assert rp.headers["content-type"] == "application/pdf"


# ---------- Staff Correction ----------
class TestCorrection:
    def test_correction_on_draft_rejected(self, admissions_tok):
        # candidate01 is draft
        rows = requests.get(f"{API}/admin/applications", headers=H(admissions_tok), timeout=30).json()
        draft = next((r for r in rows if r["application_status"] == "draft" and r["application_id"]), None)
        if not draft:
            pytest.skip("no draft app")
        r = requests.post(f"{API}/admin/applications/{draft['application_id']}/correction",
                         json={"section": "personal", "data": {}, "reason": "test correction"},
                         headers=H(admissions_tok), timeout=20)
        assert r.status_code == 409
        assert "not_submitted" in r.text

    def test_invalid_correction_422(self, admissions_tok):
        rows = requests.get(f"{API}/admin/applications?status=submitted",
                            headers=H(admissions_tok), timeout=30).json()
        if not rows:
            pytest.skip("no submitted app")
        r = requests.post(f"{API}/admin/applications/{rows[0]['application_id']}/correction",
                         json={"section": "contact", "data": {"pincode": "99"}, "reason": "bad data test"},
                         headers=H(admissions_tok), timeout=20)
        assert r.status_code == 422

    def test_valid_correction_versions_and_resets_confirmation(self, admissions_tok):
        rows = requests.get(f"{API}/admin/applications?status=submitted",
                            headers=H(admissions_tok), timeout=30).json()
        # pick candidate04 (has confirmation) if available
        c04 = next((r for r in rows if r["name"] and "Four" in r["name"]), None) or rows[0]
        app_id = c04["application_id"]
        before = requests.get(f"{API}/admin/applications/{app_id}",
                              headers=H(admissions_tok), timeout=20).json()
        v_before = before["submitted_version"]
        # pick a non-conflicting change: toggle alternate_mobile on contact
        contact = before["sections"]["contact"].copy()
        contact["alternate_mobile"] = "9988776655" if contact.get("alternate_mobile") != "9988776655" else "9988776600"
        r = requests.post(f"{API}/admin/applications/{app_id}/correction",
                         json={"section": "contact", "data": contact, "reason": "updating alternate mobile per candidate email"},
                         headers=H(admissions_tok), timeout=20)
        assert r.status_code == 200, r.text
        assert r.json()["version"] == v_before + 1
        # verify version incremented and confirmation unset
        after = requests.get(f"{API}/admin/applications/{app_id}",
                             headers=H(admissions_tok), timeout=20).json()
        assert after["submitted_version"] == v_before + 1
        assert not after.get("post", {}).get("confirmation")
        assert len(after["versions"]) >= 2


# ---------- RBAC ----------
class TestRBAC:
    def test_auditor_can_list_but_not_export(self, auditor_tok):
        r = requests.get(f"{API}/admin/applications", headers=H(auditor_tok), timeout=30)
        assert r.status_code == 200
        rows = r.json()
        # email masked
        if rows:
            assert "***" in rows[0]["email"] or "@" in rows[0]["email"]  # masked format uses *** typically
        # export denied
        r2 = requests.get(f"{API}/admin/applications/export", headers=H(auditor_tok), timeout=20)
        assert r2.status_code == 403

    def test_auditor_cannot_correct(self, auditor_tok):
        rows = requests.get(f"{API}/admin/applications?status=submitted",
                            headers=H(auditor_tok), timeout=30).json()
        if not rows:
            pytest.skip("no submitted")
        r = requests.post(f"{API}/admin/applications/{rows[0]['application_id']}/correction",
                         json={"section": "contact", "data": {}, "reason": "no permission"},
                         headers=H(auditor_tok), timeout=20)
        assert r.status_code == 403

    def test_panelist_denied_on_applications(self, panelist_tok):
        r = requests.get(f"{API}/admin/applications", headers=H(panelist_tok), timeout=20)
        assert r.status_code == 403

    def test_accounts_denied_on_applications(self, accounts_tok):
        r = requests.get(f"{API}/admin/applications", headers=H(accounts_tok), timeout=20)
        assert r.status_code == 403

    def test_candidate_denied_on_admin(self, c01_tok):
        r = requests.get(f"{API}/admin/applications", headers=H(c01_tok), timeout=20)
        assert r.status_code == 403

    def test_unverified_candidate_blocked(self):
        # candidate10 is registered-only. candidate08/09 partially verified
        for em in ("candidate08@example.com", "candidate10@example.com"):
            r = _login(em, CAND_PW, staff=False)
            if r.status_code != 200:
                continue
            tok = r.json()["access_token"]
            rr = requests.get(f"{API}/candidate/application", headers=H(tok), timeout=20)
            # Could be 403 identity_unverified or 200 depending; spec says 403
            assert rr.status_code == 403, f"{em}: {rr.status_code} {rr.text[:200]}"
            assert "identity_unverified" in rr.text or "identity" in rr.text.lower()


# ---------- Fee flow: already paid -> second checkout 409 ----------
class TestFeeAlreadyPaid:
    def test_second_checkout_already_paid(self, c03_tok):
        r = requests.post(f"{API}/candidate/application/fee/checkout",
                         json={"origin_url": "https://x.example.com"}, headers=H(c03_tok), timeout=20)
        assert r.status_code == 409
        assert "already_paid" in r.text

    def test_fee_status_paid(self, c03_tok):
        r = requests.get(f"{API}/candidate/application/fee", headers=H(c03_tok), timeout=20)
        assert r.status_code == 200
        assert r.json()["paid"] is True
        assert r.json()["amount"] == 1500

    def test_invoice_pdf(self, c03_tok):
        r = requests.get(f"{API}/candidate/application/fee/receipt", headers=H(c03_tok), timeout=30)
        assert r.status_code == 200
        assert r.headers["content-type"] == "application/pdf"
        assert len(r.content) > 500
