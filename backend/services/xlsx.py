import io

from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def xlsx_response(sheets: list[tuple[str, list[str], list[list]]], filename: str) -> Response:
    wb = Workbook()
    wb.remove(wb.active)
    for title, head, rows in sheets:
        ws = wb.create_sheet(title[:31])
        ws.append(head)
        for r in rows:
            ws.append(r)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="9E1B32")
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(40, max(12, max(len(str(x.value or "")) for x in col) + 2))
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type=XLSX, headers={"Content-Disposition": f'attachment; filename="{filename}"'})
