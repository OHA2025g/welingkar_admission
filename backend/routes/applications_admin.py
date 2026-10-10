import io
import re
from collections import Counter
from typing import Optional

from bson import ObjectId
from fastapi import APIRouter, Depends, Query, Request
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel, Field

from app_schema import SECTIONS, months, Experience
from core import db, ist_str, now_utc
from rbac import UNMASKED_CONTACT_ROLES
from routes.application import Application, render_pdf, snapshot, validate_section, write_version
from routes.config import oid
from routes.otp import _mask
from security import User, audit, err, require

router = APIRouter(prefix="/api/admin", tags=["applications-admin"])


async def _rows(viewer: User) -> list[dict]:
    apps = {a["candidate_id"]: Application.from_mongo(a) async for a in db.applications.find()}
    paid = set(await db.payment_orders.distinct("candidate_id", {"fee_code": "application_fee", "status": "paid"}))
    progs = {p["code"]: p["name"] async for p in db.programs.find({}, {"code": 1, "name": 1})}
    rows = []
    async for u in db.users.find({"role": "candidate"}).sort("created_at", 1):
        cid, a = str(u["_id"]), apps.get(str(u["_id"]))
        prefs = (a.sections.get("preferences") or {}).get("ranked", []) if a else []
        masked = viewer.role not in UNMASKED_CONTACT_ROLES
        rows.append({
            "candidate_id": cid, "candidate_ref": u.get("candidate_ref"), "name": u["name"],
            "email": _mask("email", u["email"]) if masked else u["email"],
            "mobile": _mask("whatsapp", u.get("mobile") or "+0000000000") if masked else u.get("mobile"),
            "state": u.get("state"), "city": u.get("city"), "identity_status": u.get("identity_status"),
            "application_id": a.id if a else None, "application_ref": a.application_ref if a else None,
            "application_status": a.status if a else "not_started", "fee_paid": cid in paid,
            "sections_complete": sum(v == "complete" for v in a.section_status.values()) if a else 0,
            "first_preference": progs.get(prefs[0], prefs[0]) if prefs else None, "preferences": prefs,
            "submitted_at": a.submitted_at if a else None,
            "declaration": bool(a and a.post.get("declaration")), "essays": bool(a and (a.post.get("essays") or {}).get("submitted_at")),
            "confirmation": bool(a and a.post.get("confirmation")), "app": a,
        })
    return rows


@router.get("/part-a/dashboard")
async def part_a_dashboard(viewer: User = Depends(require("applications.read"))):
    rows = await _rows(viewer)
    by_state: dict = {}
    for r in rows:
        s = by_state.setdefault(r["state"] or "—", {"state": r["state"] or "—", "registered": 0, "paid": 0, "unpaid": 0, "submitted": 0})
        s["registered"] += 1
        s["paid" if r["fee_paid"] else "unpaid"] += 1
        s["submitted"] += r["application_status"] == "submitted"
    sub = [r for r in rows if r["application_status"] == "submitted"]
    return {
        "counts": {"registered": len(rows), "verified": sum(r["identity_status"] == "verified" for r in rows),
                   "paid": sum(r["fee_paid"] for r in rows), "draft": sum(r["application_status"] == "draft" for r in rows),
                   "submitted": len(sub), "declarations": sum(r["declaration"] for r in sub),
                   "essays": sum(r["essays"] for r in sub), "confirmed": sum(r["confirmation"] for r in sub)},
        "by_state": sorted(by_state.values(), key=lambda x: -x["registered"]),
        "first_preferences": [{"program": k, "count": v} for k, v in Counter(r["first_preference"] for r in rows if r["first_preference"]).most_common()],
    }


def _filter(rows, q, status, fee, state):
    if q:
        rx = re.compile(re.escape(q), re.I)
        rows = [r for r in rows if any(rx.search(str(r.get(k) or "")) for k in ("name", "email", "candidate_ref", "application_ref"))]
    if status:
        rows = [r for r in rows if r["application_status"] == status]
    if fee:
        rows = [r for r in rows if r["fee_paid"] == (fee == "paid")]
    if state:
        rows = [r for r in rows if r["state"] == state]
    return rows


def _application_pipeline(q: Optional[str], status: Optional[str], fee: Optional[str], state: Optional[str]) -> list[dict]:
    early: dict = {"role": "candidate"}
    if state:
        early["state"] = state
    pipeline: list[dict] = [
        {"$match": early},
        {"$lookup": {"from": "applications", "let": {"cid": {"$toString": "$_id"}}, "pipeline": [
            {"$match": {"$expr": {"$eq": ["$candidate_id", "$$cid"]}}},
            {"$project": {"status": 1, "application_ref": 1, "section_status": 1, "submitted_at": 1,
                          "ranked": "$sections.preferences.ranked"}},
        ], "as": "apps"}},
        {"$lookup": {"from": "payment_orders", "let": {"cid": {"$toString": "$_id"}}, "pipeline": [
            {"$match": {"$expr": {"$and": [
                {"$eq": ["$candidate_id", "$$cid"]},
                {"$eq": ["$fee_code", "application_fee"]},
                {"$eq": ["$status", "paid"]},
            ]}}},
            {"$limit": 1}, {"$project": {"_id": 1}},
        ], "as": "paid"}},
        {"$addFields": {"app": {"$arrayElemAt": ["$apps", 0]}, "fee_paid": {"$gt": [{"$size": "$paid"}, 0]}}},
    ]
    late: dict = {}
    if q:
        rx = {"$regex": re.escape(q), "$options": "i"}
        late["$or"] = [{"name": rx}, {"email": rx}, {"candidate_ref": rx}, {"app.application_ref": rx}]
    if status == "not_started":
        late["app"] = None
    elif status:
        late["app.status"] = status
    if fee:
        late["fee_paid"] = fee == "paid"
    if late:
        pipeline.append({"$match": late})
    pipeline.append({"$sort": {"created_at": 1}})
    return pipeline


async def _page_applications(viewer: User, q: Optional[str], status: Optional[str], fee: Optional[str],
                             state: Optional[str], page: int, page_size: int) -> dict:
    size = min(max(page_size, 1), 200)
    start = (max(page, 1) - 1) * size
    pipeline = _application_pipeline(q, status, fee, state)
    pipeline.append({"$facet": {"rows": [{"$skip": start}, {"$limit": size}], "meta": [{"$count": "total"}]}})
    packed = await db.users.aggregate(pipeline).to_list(1)
    facet = packed[0] if packed else {"rows": [], "meta": []}
    docs = facet.get("rows") or []
    total = facet["meta"][0]["total"] if facet.get("meta") else 0
    codes = {((d.get("app") or {}).get("ranked") or [None])[0] for d in docs}
    codes.discard(None)
    progs = {p["code"]: p["name"] async for p in db.programs.find({"code": {"$in": list(codes)}}, {"code": 1, "name": 1})} if codes else {}
    masked = viewer.role not in UNMASKED_CONTACT_ROLES
    rows = []
    for d in docs:
        app = d.get("app")
        prefs = (app or {}).get("ranked") or []
        first = prefs[0] if prefs else None
        mobile = d.get("mobile") or ""
        rows.append({
            "candidate_id": str(d["_id"]), "candidate_ref": d.get("candidate_ref"), "name": d["name"],
            "email": _mask("email", d["email"]) if masked else d["email"],
            "mobile": _mask("whatsapp", mobile) if masked and mobile else mobile or None,
            "state": d.get("state"), "city": d.get("city"), "identity_status": d.get("identity_status"),
            "application_id": str(app["_id"]) if app else None,
            "application_ref": app.get("application_ref") if app else None,
            "application_status": app.get("status") if app else "not_started",
            "fee_paid": bool(d.get("fee_paid")),
            "sections_complete": sum(v == "complete" for v in (app or {}).get("section_status", {}).values()) if app else 0,
            "first_preference": progs.get(first, first) if first else None, "preferences": prefs,
            "submitted_at": app.get("submitted_at") if app else None,
        })
    return {"rows": rows, "total": total, "page": max(page, 1), "page_size": size}


@router.get("/applications")
async def list_applications(q: Optional[str] = None, status: Optional[str] = None, fee: Optional[str] = None,
                            state: Optional[str] = None, page: Optional[int] = None,
                            page_size: int = Query(50, ge=1, le=200),
                            viewer: User = Depends(require("applications.read"))):
    if page is None:
        return [{k: v for k, v in r.items() if k != "app"} for r in _filter(await _rows(viewer), q, status, fee, state)]
    return await _page_applications(viewer, q, status, fee, state, page, page_size)


@router.get("/applications/export")
async def export(request: Request, q: Optional[str] = None, status: Optional[str] = None, fee: Optional[str] = None,
                 state: Optional[str] = None, viewer: User = Depends(require("applications.export"))):
    rows = _filter(await _rows(viewer), q, status, fee, state)
    wb = Workbook()
    ws = wb.active
    ws.title = "Candidates"
    head = ["Candidate ID", "Application ref", "Name", "Email", "Mobile", "State", "City", "Identity", "Fee paid",
            "Status", "Submitted (IST)", "DOB", "Gender", "Category", "Father", "Mother", "Family income (INR)",
            "Pincode", "X %", "XII %", "Graduation", "Grad score", "Work months", "Exams", "Preferences",
            "Declaration", "Essays submitted", "Data confirmed"]
    ws.append(head)
    for r in rows:
        s = r["app"].sections if r["app"] else {}
        p, f, c, e = s.get("personal") or {}, s.get("family") or {}, s.get("contact") or {}, s.get("education") or {}
        g = e.get("graduation") or {}
        w = s.get("work") or {}
        try:
            wm = sum(months(Experience(**x)) for x in w.get("experiences", [])) if w.get("has_experience") else 0
        except Exception:
            wm = ""
        exams = "; ".join(f"{x.get('exam')} {x.get('year')}: {x.get('percentile') or x.get('score') or x.get('status')}"
                          for x in (s.get("exams") or {}).get("exams", []))
        ws.append([r["candidate_ref"], r["application_ref"], r["name"], r["email"], r["mobile"], r["state"], r["city"],
                   r["identity_status"], "Yes" if r["fee_paid"] else "No", r["application_status"],
                   ist_str(r["submitted_at"]) if r["submitted_at"] else "", p.get("dob"), p.get("gender"), p.get("category"),
                   f.get("father_name"), f.get("mother_name"), f.get("annual_family_income_inr"), c.get("pincode"),
                   (e.get("class_x") or {}).get("percentage"), (e.get("class_xii") or {}).get("percentage"),
                   f"{g.get('degree', '')} {g.get('discipline', '')}".strip(), g.get("score"), wm, exams,
                   " > ".join(r["preferences"]), "Yes" if r["declaration"] else "No", "Yes" if r["essays"] else "No",
                   "Yes" if r["confirmation"] else "No"])
    st = wb.create_sheet("State-wise paid")
    st.append(["State", "Registered", "Paid", "Unpaid", "Submitted"])
    for s in (await part_a_dashboard(viewer))["by_state"]:
        st.append([s["state"], s["registered"], s["paid"], s["unpaid"], s["submitted"]])
    for sheet in (ws, st):
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="9E1B32")
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = min(40, max(12, max(len(str(x.value or "")) for x in col) + 2))
    buf = io.BytesIO()
    wb.save(buf)
    await audit(request, viewer, "applications.exported", "export", None,
                after={"rows": len(rows), "filters": {"q": q, "status": status, "fee": fee, "state": state}})
    name = f"weschool-part-a-{now_utc():%Y%m%d-%H%M}.xlsx"
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{name}"'})


@router.get("/applications/{app_id}")
async def get_application(app_id: str, viewer: User = Depends(require("applications.read"))):
    a = Application.from_mongo(await db.applications.find_one({"_id": oid(app_id)}))
    if not a:
        raise err(404, "not_found", "Application not found")
    u = await db.users.find_one({"_id": ObjectId(a.candidate_id)})
    masked = viewer.role not in UNMASKED_CONTACT_ROLES
    versions = [{"version": v["version"], "reason": v["reason"], "created_by": v["created_by"], "created_at": v["created_at"]}
                async for v in db.application_versions.find({"application_id": a.id}).sort("version", 1)]
    paid = await db.payment_orders.find_one({"candidate_id": a.candidate_id, "fee_code": "application_fee", "status": "paid"})
    return {**a.model_dump(), "versions": versions, "fee_paid": bool(paid), "fee_order_ref": paid["order_ref"] if paid else None,
            "sections_meta": [{"key": k, "title": t} for k, (t, _) in SECTIONS.items()],
            "candidate": {"name": u["name"], "candidate_ref": u.get("candidate_ref"), "state": u.get("state"), "city": u.get("city"),
                          "email": _mask("email", u["email"]) if masked else u["email"],
                          "mobile": _mask("whatsapp", u.get("mobile") or "+0000000000") if masked else u.get("mobile"),
                          "identity_status": u.get("identity_status")}}


@router.get("/applications/{app_id}/pdf")
async def application_pdf(app_id: str, _: User = Depends(require("applications.read"))):
    data, name = await render_pdf(str(oid(app_id)))
    return Response(data, media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="{name}"'})


class CorrectionIn(BaseModel):
    section: str
    data: dict
    reason: str = Field(min_length=5, max_length=500)


@router.post("/applications/{app_id}/correction")
async def correct(app_id: str, body: CorrectionIn, request: Request, actor: User = Depends(require("applications.correct"))):
    _id = oid(app_id)
    a = Application.from_mongo(await db.applications.find_one({"_id": _id}))
    if not a:
        raise err(404, "not_found", "Application not found")
    if a.status != "submitted":
        raise err(409, "not_submitted", "Only submitted applications can be corrected")
    if body.section not in SECTIONS:
        raise err(422, "unknown_section", "Unknown section")
    clean, errors = await validate_section(body.section, body.data, a.cycle_id)
    if errors:
        raise err(422, "invalid", errors[0]["message"], errors=errors)
    before = a.sections.get(body.section)
    if before == clean:
        raise err(409, "unchanged", "No changes to save")
    new_v = a.submitted_version + 1
    res = await db.applications.update_one({"_id": _id, "submitted_version": a.submitted_version}, {
        "$set": {f"sections.{body.section}": clean, "submitted_version": new_v, "updated_at": now_utc()},
        "$unset": {"post.confirmation": ""}})
    if not res.modified_count:
        raise err(409, "stale_version", "Application changed meanwhile. Reload and try again.")
    a = Application.from_mongo(await db.applications.find_one({"_id": _id}))
    await write_version(a.id, new_v, await snapshot(a, await db.users.find_one({"_id": ObjectId(a.candidate_id)})),
                        f"correction: {body.reason}", actor.email)
    await audit(request, actor, "application.corrected", "application", a.id, before=before, after=clean,
                reason=f"[{body.section}] v{new_v - 1}->v{new_v}: {body.reason}")
    return {"version": new_v, "section": body.section}
