import csv
import hashlib
import io
from datetime import datetime
from typing import Literal, Optional

from bson import ObjectId
from fastapi import APIRouter, Depends, File, Form, Request, UploadFile
from openpyxl import load_workbook
from pydantic import BaseModel, Field
from pymongo.errors import BulkWriteError

from core import BaseDocument, db, now_utc
from routes.application import active_cycle
from security import User, audit, err, require
from services.notify import notify
from services.xlsx import xlsx_response

router = APIRouter(prefix="/api", tags=["partb"])
COLUMNS = ["cycle_code", "program_code", "application_reference", "outcome", "merit_rank", "waitlist_rank",
           "choice_position", "list_version"]
OUTCOMES = ("selected", "waitlisted", "rejected")
MAX_ROWS = 5000


class ImportBatch(BaseDocument):
    cycle_id: str
    kind: str
    filename: str
    checksum: str
    mode: str
    status: str
    list_version: Optional[int] = None
    rows_total: int
    rows_valid: int
    errors: list[dict] = []
    rows: list[dict] = []
    created_by: str
    created_at: datetime = Field(default_factory=now_utc)
    committed_at: Optional[datetime] = None
    committed_by: Optional[str] = None
    inserted: int = 0
    rolled_back_at: Optional[datetime] = None
    rolled_back_by: Optional[str] = None
    rollback_reason: Optional[str] = None


class Eligibility(BaseDocument):
    cycle_id: str
    candidate_id: str
    application_id: str
    application_ref: str
    program_code: str
    outcome: str
    merit_rank: Optional[int] = None
    waitlist_rank: Optional[int] = None
    choice_position: Optional[int] = None
    import_id: str
    status: str = "active"
    created_at: datetime = Field(default_factory=now_utc)


class Consent(BaseDocument):
    cycle_id: str
    candidate_id: str
    kind: Literal["initial", "betterment"] = "initial"
    batch_id: Optional[str] = None
    courses: dict[str, str]
    overall: bool = False
    version: int = 1
    status: str = "active"
    actor: str
    reason: Optional[str] = None
    created_at: datetime = Field(default_factory=now_utc)


async def eligibility_for(cid: str) -> list[Eligibility]:
    return [Eligibility.from_mongo(e) async for e in db.partb_eligibility.find({"candidate_id": cid, "status": "active"}).sort("choice_position", 1)]


def consentable(rows: list[Eligibility]) -> list[str]:
    return [e.program_code for e in rows if e.outcome in ("selected", "waitlisted")]


async def active_consent(cid: str) -> Optional[Consent]:
    return Consent.from_mongo(await db.partb_consents.find_one({"candidate_id": cid, "kind": "initial", "status": "active"}))


def _int(v) -> Optional[int]:
    if v in (None, ""):
        return None
    f = float(v)
    if f != int(f) or f < 1:
        raise ValueError
    return int(f)


def parse_file(name: str, data: bytes) -> list[dict]:
    low = name.lower()
    if low.endswith(".csv"):
        try:
            rows = list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))
        except UnicodeDecodeError:
            raise err(422, "bad_encoding", "CSV must be UTF-8 encoded")
    elif low.endswith(".xlsx"):
        try:
            ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        except Exception:
            raise err(422, "bad_file", "Could not read the Excel file")
        it = ws.iter_rows(values_only=True)
        head = [str(h or "").strip().lower() for h in next(it, [])]
        rows = [dict(zip(head, r)) for r in it if any(c not in (None, "") for c in r)]
    else:
        raise err(422, "bad_type", "Upload a .csv or .xlsx file")
    rows = [{str(k).strip().lower(): (str(v).strip() if v is not None else "") for k, v in r.items() if k} for r in rows]
    if not rows:
        raise err(422, "empty", "The file has no data rows")
    missing = [c for c in ("cycle_code", "program_code", "application_reference", "outcome") if c not in rows[0]]
    if missing:
        raise err(422, "missing_columns", f"Missing columns: {', '.join(missing)}", expected=COLUMNS)
    if len(rows) > MAX_ROWS:
        raise err(422, "too_many_rows", f"At most {MAX_ROWS} rows per file")
    return rows


async def validate_rows(raw: list[dict]) -> tuple[list[dict], list[dict]]:
    c = await active_cycle()
    progs = {p["code"] async for p in db.programs.find({"cycle_id": c.id}, {"code": 1})}
    apps = {a["application_ref"]: a async for a in db.applications.find({"cycle_id": c.id, "status": "submitted"},
                                                                         {"application_ref": 1, "candidate_id": 1})}
    approved = {v["version"] async for v in db.result_versions.find({"cycle_id": c.id, "status": {"$in": ["approved", "published", "superseded"]},
                                                                     "approved_at": {"$ne": None}}, {"version": 1})}
    pairs, ranks = set(), set()
    async for e in db.partb_eligibility.find({"cycle_id": c.id, "status": "active"}):
        pairs.add((e["application_ref"], e["program_code"]))
        if e.get("merit_rank"):
            ranks.add((e["program_code"], "merit", e["merit_rank"]))
        if e.get("waitlist_rank"):
            ranks.add((e["program_code"], "waitlist", e["waitlist_rank"]))
    valid, errors = [], []
    for n, r in enumerate(raw, start=2):
        errs = []

        def bad(field, msg):
            errs.append({"row": n, "field": field, "message": msg})
        code, ref, outcome = r.get("program_code", "").upper(), r.get("application_reference", ""), r.get("outcome", "").lower()
        if r.get("cycle_code") != c.code:
            bad("cycle_code", f"Cycle must be {c.code}")
        if code not in progs:
            bad("program_code", f"Unknown program '{code}'")
        app = apps.get(ref)
        if not app:
            bad("application_reference", f"No submitted application '{ref}'")
        if outcome not in OUTCOMES:
            bad("outcome", "Outcome must be selected, waitlisted or rejected")
        vals = {}
        for f in ("merit_rank", "waitlist_rank", "choice_position", "list_version"):
            try:
                vals[f] = _int(r.get(f))
            except ValueError:
                bad(f, f"{f.replace('_', ' ').capitalize()} must be a positive whole number")
        if outcome == "selected" and not vals.get("merit_rank") and "merit_rank" not in [e["field"] for e in errs]:
            bad("merit_rank", "Selected rows need a merit rank")
        if outcome == "waitlisted" and not vals.get("waitlist_rank") and "waitlist_rank" not in [e["field"] for e in errs]:
            bad("waitlist_rank", "Waitlisted rows need a waitlist rank")
        if outcome == "rejected" and (vals.get("merit_rank") or vals.get("waitlist_rank")):
            bad("outcome", "Rejected rows cannot carry a rank")
        if vals.get("list_version") and vals["list_version"] not in approved:
            bad("list_version", f"Result version {vals['list_version']} is not approved")
        if (ref, code) in pairs:
            bad("application_reference", f"Duplicate candidate-course row for {ref} / {code}")
        for kind in ("merit", "waitlist"):
            rk = vals.get(f"{kind}_rank")
            if rk and (code, kind, rk) in ranks:
                bad(f"{kind}_rank", f"{kind.capitalize()} rank {rk} is already used for {code}")
        if errs:
            errors.extend(errs)
            continue
        pairs.add((ref, code))
        for kind in ("merit", "waitlist"):
            if vals.get(f"{kind}_rank"):
                ranks.add((code, kind, vals[f"{kind}_rank"]))
        valid.append({"cycle_id": c.id, "candidate_id": app["candidate_id"], "application_id": str(app["_id"]), "application_ref": ref,
                      "program_code": code, "outcome": outcome, "merit_rank": vals.get("merit_rank"),
                      "waitlist_rank": vals.get("waitlist_rank"), "choice_position": vals.get("choice_position")})
    return valid, errors


async def insert_rows(batch_id: str, rows: list[dict]) -> int:
    if not rows:
        return 0
    docs = [Eligibility(import_id=batch_id, **r).to_mongo() for r in rows]
    try:
        return len((await db.partb_eligibility.insert_many(docs, ordered=False)).inserted_ids)
    except BulkWriteError as e:
        return e.details["nInserted"]


async def run_import(request: Request, user: User, kind: str, filename: str, checksum: str, raw: list[dict], mode: str,
                     list_version: Optional[int] = None) -> dict:
    c = await active_cycle()
    if await db.partb_import_batches.find_one({"cycle_id": c.id, "checksum": checksum, "status": "inserted"}):
        raise err(409, "duplicate_file", "This exact list is already inserted. Roll it back first to re-import.")
    valid, errors = await validate_rows(raw)
    b = ImportBatch(cycle_id=c.id, kind=kind, filename=filename, checksum=checksum, mode=mode, list_version=list_version,
                    status="validated" if mode == "validate" else "inserted", rows_total=len(raw), rows_valid=len(valid),
                    errors=errors, rows=valid if mode == "validate" else [], created_by=user.email)
    b.id = str((await db.partb_import_batches.insert_one(b.to_mongo())).inserted_id)
    if mode == "insert":
        b.inserted = await insert_rows(b.id, valid)
        await db.partb_import_batches.update_one({"_id": ObjectId(b.id)}, {"$set": {
            "inserted": b.inserted, "committed_at": now_utc(), "committed_by": user.email}})
        await _notify_new(valid)
    await audit(request, user, f"partb.import_{b.status}", "partb_import", b.id,
                after={"filename": filename, "rows": len(raw), "valid": len(valid), "errors": len(errors), "mode": mode})
    return b.model_dump(exclude={"rows"})


async def _notify_new(rows: list[dict]):
    for cid in {r["candidate_id"] for r in rows if r["outcome"] != "rejected"}:
        await notify(cid, "partb_open", "Part B is open for you", "Your merit/waitlist status is available. Review it and give your consent.")


@router.post("/partb/imports")
async def import_file(request: Request, file: UploadFile = File(...), mode: Literal["validate", "insert"] = Form("validate"),
                      user: User = Depends(require("partb.import"))):
    data = await file.read(2_000_001)
    if len(data) > 2_000_000:
        raise err(422, "too_large", "File must be 2 MB or smaller")
    raw = parse_file(file.filename or "upload", data)
    return await run_import(request, user, "file", file.filename, hashlib.sha256(data).hexdigest(), raw, mode)


class FromResultIn(BaseModel):
    version: int
    mode: Literal["validate", "insert"] = "validate"


@router.post("/partb/imports/from-result")
async def import_from_result(body: FromResultIn, request: Request, user: User = Depends(require("partb.import"))):
    c = await active_cycle()
    rv = await db.result_versions.find_one({"cycle_id": c.id, "version": body.version})
    if not rv:
        raise err(404, "not_found", "Result version not found")
    if rv["status"] != "published":
        raise err(409, "not_published", "Only the published result version can be inserted into Part B")
    raw, wl = [], {}
    for r in sorted(rv["rows"], key=lambda x: (x["program_code"], x["rank"] or 10 ** 6)):
        out = {"selected": "selected", "waitlisted": "waitlisted", "not_selected": "rejected"}.get(r["outcome"])
        if not out:
            continue
        w = None
        if out == "waitlisted":
            wl[r["program_code"]] = w = wl.get(r["program_code"], 0) + 1
        raw.append({"cycle_code": c.code, "program_code": r["program_code"], "application_reference": r["application_ref"],
                    "outcome": out, "merit_rank": str(r["rank"]) if out == "selected" else "", "waitlist_rank": str(w or ""),
                    "choice_position": str(r["choice"]), "list_version": str(body.version)})
    if not raw:
        raise err(422, "empty", "The result version has no selected or waitlisted rows")
    checksum = hashlib.sha256(f"result-v{body.version}-{rv['inputs']['checksum']}".encode()).hexdigest()
    return await run_import(request, user, "result_version", f"Result v{body.version}", checksum, raw, body.mode, body.version)


@router.get("/partb/imports")
async def list_imports(_: User = Depends(require("partb.import"))):
    c = await active_cycle()
    return [ImportBatch.from_mongo(b).model_dump(exclude={"rows", "errors"}) | {"error_count": len(b.get("errors", []))}
            async for b in db.partb_import_batches.find({"cycle_id": c.id}).sort("created_at", -1)]


async def _batch(bid: str) -> ImportBatch:
    b = ImportBatch.from_mongo(await db.partb_import_batches.find_one({"_id": ObjectId(bid)})) if ObjectId.is_valid(bid) else None
    if not b:
        raise err(404, "not_found", "Import batch not found")
    return b


@router.get("/partb/imports/{bid}")
async def get_import(bid: str, _: User = Depends(require("partb.import"))):
    b = await _batch(bid)
    return b.model_dump() | {"eligibility_rows": await db.partb_eligibility.count_documents({"import_id": b.id, "status": "active"})}


@router.get("/partb/imports/{bid}/error-report")
async def error_report(bid: str, _: User = Depends(require("partb.import"))):
    b = await _batch(bid)
    return xlsx_response([("Row errors", ["Row", "Field", "Message"], [[e["row"], e["field"], e["message"]] for e in b.errors])],
                         f"partb-import-errors-{b.id[-6:]}.xlsx")


@router.post("/partb/imports/{bid}/commit")
async def commit_import(bid: str, request: Request, user: User = Depends(require("partb.import"))):
    b = await _batch(bid)
    res = await db.partb_import_batches.update_one({"_id": ObjectId(b.id), "status": "validated"}, {"$set": {"status": "committing"}})
    if not res.modified_count:
        raise err(409, "not_validated", "Only a validated batch can be inserted")
    raw = [{"cycle_code": (await active_cycle()).code, "program_code": r["program_code"], "application_reference": r["application_ref"],
            "outcome": r["outcome"], **{k: str(r[k] or "") for k in ("merit_rank", "waitlist_rank", "choice_position")},
            "list_version": str(b.list_version or "")} for r in b.rows]
    valid, errors = await validate_rows(raw)
    inserted = await insert_rows(b.id, valid)
    await db.partb_import_batches.update_one({"_id": ObjectId(b.id)}, {"$set": {
        "status": "inserted", "inserted": inserted, "committed_at": now_utc(), "committed_by": user.email,
        "errors": b.errors + [{**e, "message": f"[at insert] {e['message']}"} for e in errors]}})
    await _notify_new(valid)
    await audit(request, user, "partb.import_inserted", "partb_import", b.id, after={"inserted": inserted, "late_errors": len(errors)})
    return (await _batch(bid)).model_dump(exclude={"rows"})


class RollbackIn(BaseModel):
    reason: str = Field(min_length=5, max_length=300)


@router.post("/partb/imports/{bid}/rollback")
async def rollback_import(bid: str, body: RollbackIn, request: Request, user: User = Depends(require("partb.import"))):
    b = await _batch(bid)
    if b.status != "inserted":
        raise err(409, "not_inserted", "Only an inserted batch can be rolled back")
    pairs = [(e["candidate_id"], e["program_code"]) async for e in db.partb_eligibility.find({"import_id": b.id, "status": "active"})]
    for cid, code in pairs:
        if await db.partb_allotments.find_one({"candidate_id": cid, "program_code": code, "status": {"$ne": "rejected"}}):
            raise err(409, "has_allotments", "Candidates from this batch already have allotments. Cancel them before rolling back.")
    res = await db.partb_import_batches.update_one({"_id": ObjectId(b.id), "status": "inserted"}, {"$set": {
        "status": "rolled_back", "rolled_back_at": now_utc(), "rolled_back_by": user.email, "rollback_reason": body.reason}})
    if not res.modified_count:
        raise err(409, "not_inserted", "Batch already rolled back")
    n = (await db.partb_eligibility.update_many({"import_id": b.id, "status": "active"}, {"$set": {"status": "rolled_back"}})).modified_count
    await audit(request, user, "partb.import_rolled_back", "partb_import", b.id, after={"rows": n}, reason=body.reason)
    return (await _batch(bid)).model_dump(exclude={"rows"})


# ---------- consent ----------
class ConsentIn(BaseModel):
    courses: dict[str, Literal["accept", "decline"]]
    overall: bool


def check_consent(body: ConsentIn, eligible: list[str]):
    if set(body.courses) - set(eligible):
        raise err(422, "ineligible_course", "Consent can only cover courses you are eligible for")
    if set(eligible) - set(body.courses):
        raise err(422, "incomplete", "Respond to every eligible course")
    if not body.overall:
        raise err(422, "overall_required", "Overall admission consent is required to continue to documents")
    if "accept" not in body.courses.values():
        raise err(422, "no_course", "Accept at least one course to continue")


@router.post("/candidate/partb/consent")
async def give_consent(body: ConsentIn, request: Request, user: User = Depends(require("self.application"))):
    rows = await eligibility_for(user.id)
    eligible = consentable(rows)
    if not eligible:
        raise err(403, "not_eligible", "You are not eligible for Part B")
    if await active_consent(user.id):
        raise err(409, "already_consented", "Consent already recorded. Contact the admissions office to change it.")
    check_consent(body, eligible)
    c = Consent(cycle_id=rows[0].cycle_id, candidate_id=user.id, courses=body.courses, overall=True, actor=user.email)
    c.id = str((await db.partb_consents.insert_one(c.to_mongo())).inserted_id)
    await audit(request, user, "partb.consent_given", "partb_consent", c.id, after=body.model_dump())
    await notify(user.id, "consent", "Consent recorded", "Your course and admission consent is recorded. Upload your documents next.")
    return c.model_dump()


class CorrectConsentIn(ConsentIn):
    reason: str = Field(min_length=5, max_length=300)


@router.post("/partb/consents/{cid}/correct")
async def correct_consent(cid: str, body: CorrectConsentIn, request: Request, user: User = Depends(require("consent.correct"))):
    rows = await eligibility_for(cid)
    check_consent(body, consentable(rows))
    held = await db.partb_allotments.find_one({"candidate_id": cid, "seat_held": True})
    if held and body.courses.get(held["program_code"]) != "accept":
        raise err(409, "seat_held", "The candidate holds a seat in this course. Cancel the allotment before declining it.")
    cur = await active_consent(cid)
    new = Consent(cycle_id=rows[0].cycle_id, candidate_id=cid, courses=body.courses, overall=True,
                  version=(cur.version + 1) if cur else 1, actor=user.email, reason=body.reason)
    if cur:
        res = await db.partb_consents.update_one({"_id": ObjectId(cur.id), "status": "active"}, {"$set": {"status": "superseded"}})
        if not res.modified_count:
            raise err(409, "stale", "Consent changed meanwhile. Reload and try again.")
    new.id = str((await db.partb_consents.insert_one(new.to_mongo())).inserted_id)
    await audit(request, user, "partb.consent_corrected", "partb_consent", new.id,
                before=cur.model_dump(include={"courses", "overall", "version"}) if cur else None,
                after=new.model_dump(include={"courses", "overall", "version"}), reason=body.reason)
    await notify(cid, "consent", "Your consent was updated", f"The admissions office updated your consent (version {new.version}).")
    return new.model_dump()


@router.get("/candidate/notifications")
async def my_notifications(user: User = Depends(require("self.application"))):
    rows = [{**{k: v for k, v in n.items() if k != "_id"}, "id": str(n["_id"])}
            async for n in db.notifications.find({"user_id": user.id}).sort("created_at", -1).limit(50)]
    return {"unread": sum(1 for n in rows if not n["read_at"]), "items": rows}


@router.post("/candidate/notifications/read-all")
async def read_all(user: User = Depends(require("self.application"))):
    await db.notifications.update_many({"user_id": user.id, "read_at": None}, {"$set": {"read_at": now_utc()}})
    return await my_notifications(user)
